import csv
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Optional

TEXT_EXTENSIONS = {
    ".txt", ".md", ".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".css",
    ".json", ".yaml", ".yml", ".xml", ".csv", ".sql", ".sh", ".java", ".go",
    ".rs", ".php", ".rb", ".c", ".cpp", ".h", ".env.example", ".toml"
}

def safe_filename(name: str) -> str:
    name = Path(name.strip() or "result.txt").name
    name = re.sub(r"[^A-Za-zА-Яа-яЁё0-9._ -]", "_", name)
    return name[:100] or "result.txt"

def make_text_file(filename: str, content: str) -> tuple[str, bytes]:
    filename = safe_filename(filename)
    return filename, content.encode("utf-8")

def make_zip(files: dict[str, str]) -> tuple[str, bytes]:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(safe_filename(name), content)
    return "project.zip", buffer.getvalue()

def read_upload(filename: str, data: bytes) -> Optional[str]:
    ext = Path(filename.lower()).suffix
    try:
        if ext in TEXT_EXTENSIONS or ext in {".txt", ".log"}:
            return data.decode("utf-8", errors="replace")[:30000]
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)[:30000]
        if ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)[:30000]
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:5]:
                lines.append(f"[Лист: {ws.title}]")
                for row in ws.iter_rows(max_row=100, values_only=True):
                    lines.append("\t".join("" if v is None else str(v) for v in row))
            return "\n".join(lines)[:30000]
        if ext == ".csv":
            return data.decode("utf-8-sig", errors="replace")[:30000]
    except Exception as exc:
        return f"Не удалось извлечь текст из файла {filename}: {exc}"
    return None

def build_document(filename: str, content: str) -> tuple[str, bytes, str]:
    filename = safe_filename(filename)
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfbase import pdfmetrics
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        buffer = io.BytesIO()
        styles = getSampleStyleSheet()
        story = []
        for line in content.splitlines() or [""]:
            escaped = (line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") or " ")
            story.append(Paragraph(escaped, styles["BodyText"]))
            story.append(Spacer(1, 2 * mm))
        SimpleDocTemplate(buffer, pagesize=A4).build(story)
        return filename, buffer.getvalue(), "application/pdf"
    if ext == ".docx":
        from docx import Document
        doc = Document()
        for line in content.splitlines():
            doc.add_paragraph(line)
        buffer = io.BytesIO()
        doc.save(buffer)
        return filename, buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if ext == ".xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for line in content.splitlines():
            ws.append(line.split("\t"))
        buffer = io.BytesIO()
        wb.save(buffer)
        return filename, buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return filename, content.encode("utf-8"), "text/plain"
